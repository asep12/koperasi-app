"""
Ekstrak Excel "PEMBUKUAN KOPERASI <tahun>.xlsx" → migrasi_data.js untuk menu
Pengaturan → Import Data Awal → "Migrasi dari Excel pembukuan".

Pakai:
    pip install openpyxl
    python tools/ekstrak_excel.py "referensi/PEMBUKUAN KOPERASI 2026.xlsx"
    python tools/ekstrak_excel.py FILE.xlsx --koreksi tools/koreksi.json --jasa-pinjaman 1.5

Yang dibaca (template Excel "PEMBUKUAN KOPERASI"):
  DATA_Anggota        daftar anggota PNS (kolom H) & Non PNS (kolom K), baris 15–64
  Simp_TH_Sblmnya     saldo pokok/wajib/sukarela per 31 Desember tahun lalu
  Piut_Th_Sblmnya     sisa piutang per 31 Desember tahun lalu
  Neraca_TH_Sblmnya   kas, inventaris, penyusutan, cadangan, dana-dana, SHU tahun lalu
  Isim(n) / Abl(n)    setoran & pengambilan simpanan bulan n
  Ang (n) / Ipin (n)  angsuran (pokok, jasa) & pinjaman baru bulan n
  KAS (n)             pos kas lain (biaya, THR, honor, pengeluaran dana, pembagian SHU)
  Simp (n) / Pinj (n) saldo akhir bulan — dipakai untuk REKONSILIASI

Setiap bulan dihitung ulang (saldo tiap anggota, piutang, kas) dan dibandingkan
dengan angka Excel. Ada selisih → file TIDAK dibuat.

File keluaran berisi nama & saldo anggota: JANGAN di-commit ke Git (sudah ada di
.gitignore) dan hapus dari Apps Script setelah migrasi selesai.

--koreksi: JSON untuk membetulkan label yang salah di sheet KAS, mis.
  [{"bulan": 1, "akun_dari": "212", "akun_ke": "211", "nominal": 50000000}]
"""
import argparse, json, os, re, sys

try:
    import openpyxl
except ImportError:
    sys.exit('Pasang dulu: pip install openpyxl')

p = argparse.ArgumentParser(description='Ekstrak Excel pembukuan koperasi → migrasi_data.js')
p.add_argument('excel')
p.add_argument('--keluar', default=os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'migrasi_data.js'))
p.add_argument('--koreksi', help='JSON koreksi akun pos kas (lihat keterangan di atas)')
p.add_argument('--bulan', type=int, help='bulan terakhir yang dimigrasi (bawaan: terakhir yang terisi)')
p.add_argument('--jasa-pinjaman', type=float, default=1.5, help='persen/bulan, menurun (bawaan 1.5)')
p.add_argument('--jasa-sukarela', type=float, default=1, help='persen (bawaan 1)')
p.add_argument('--metode-jasa', default='terendah', choices=['terendah', 'harian'],
               help='rumus jasa sukarela tahun berjalan (bawaan: saldo terendah bulan itu × %%)')
p.add_argument('--pokok', type=float, default=10000)
p.add_argument('--wajib', type=float, default=50000)
p.add_argument('--minimal-sukarela', type=float, default=10000)
arg = p.parse_args()

wb = openpyxl.load_workbook(arg.excel, read_only=True, data_only=True)
r2 = lambda x: round(float(x) + 0.0, 2) if isinstance(x, (int, float)) else 0.0
rapi = lambda t: ' '.join(str(t).split()) if t is not None else ''

# ---------------- tahun buku ----------------
tahun = None
for r in wb['KAS (1)'].iter_rows(min_row=5, max_row=9, values_only=True):
    for v in r:
        if isinstance(v, (int, float)) and 2000 < v < 2100: tahun = int(v)
if not tahun: sys.exit('Tahun buku tidak ditemukan di sheet KAS (1).')

# ---------------- anggota ----------------
anggota = []
for i, r in enumerate(wb['DATA_Anggota'].iter_rows(min_row=15, max_row=64, values_only=True)):
    for grup, kn in (('P', 7), ('N', 10)):
        nama = r[kn]
        if isinstance(nama, str) and nama.strip() and nama.strip() != '0':
            anggota.append({'kode': grup + str(i + 1), 'grup': 'PNS' if grup == 'P' else 'Non PNS', 'no': i + 1, 'nama': rapi(nama)})
KODE = {a['kode']: a for a in anggota}
if not anggota: sys.exit('DATA_Anggota kosong.')

def per_anggota(ws, baris0, kolP, kolN, nilaiP, nilaiN, n=50):
    """{kode: [nilai...]} — posisi baris sama dengan DATA_Anggota; nama dicocokkan."""
    out = {}
    rows = list(ws.iter_rows(min_row=baris0, max_row=baris0 + n - 1, values_only=True))
    for grup, kn, kv in (('P', kolP, nilaiP), ('N', kolN, nilaiN)):
        for i, r in enumerate(rows):
            kode, nama, nilai = grup + str(i + 1), r[kn], [r2(r[k]) for k in kv]
            if kode not in KODE:
                if any(abs(v) > 0.001 for v in nilai):
                    sys.exit(f'{ws.title}: baris {kode} bernilai tetapi tidak ada di DATA_Anggota: {nama} {nilai}')
                continue
            if isinstance(nama, str) and rapi(nama) != KODE[kode]['nama']:
                sys.exit(f'{ws.title}: nama tidak cocok di {kode}: {nama!r} vs DATA_Anggota {KODE[kode]["nama"]!r}')
            out[kode] = nilai
    return out

def total_bulan(ws):
    for r in ws.iter_rows(values_only=True):
        if any(isinstance(v, str) and v.strip().upper().startswith('J U M L A H') for v in r if v):
            return sum(abs(v) for v in r if isinstance(v, (int, float)))
    return 0

# ---------------- saldo 31 Des tahun lalu ----------------
simp_awal = per_anggota(wb['Simp_TH_Sblmnya'], 13, 4, 11, [5, 6, 7], [12, 13, 14])
piut_awal = per_anggota(wb['Piut_Th_Sblmnya'], 10, 6, 10, [7], [11])
ner = {}
for r in wb['Neraca_TH_Sblmnya'].iter_rows(min_row=10, max_row=40, values_only=True):
    if isinstance(r[7], str):
        ner[rapi(r[7]) + (str(r[8]) if r[8] else '')] = (r2(r[9]), r2(r[10]))
PETA_NERACA = {'KAS': ('kas', 0), 'Inventaris': ('inventaris', 0), 'Penyusutan Inventaris': ('penyusutan', 0),
               'Dana Cadangan': ('cadangan', 1), 'Dana Anggota': ('dana_anggota', 1), 'Dana Pengurus': ('dana_pengurus', 1),
               'Dana Kesej. Pegawai': ('dana_kesejahteraan', 1), 'Dana Pendidikan': ('dana_pendidikan', 1),
               'Dana Sosial': ('dana_sosial', 1), 'Dana Pemb. Daerah': ('dana_pembangunan', 1),
               'SHU Tahun' + str(tahun - 1): ('shu', 1)}
koperasi = {}
for label, (jenis, sisi) in PETA_NERACA.items():
    if label in ner:
        v = ner[label][sisi]
        koperasi[jenis] = -v if jenis == 'penyusutan' else v
if 'kas' not in koperasi: sys.exit('Neraca_TH_Sblmnya: baris KAS tidak ditemukan.')

# ---------------- bulan terisi ----------------
bulan_terisi = arg.bulan or max([b for b in range(1, 13) if total_bulan(wb[f'Isim({b})']) + total_bulan(wb[f'Ang ({b})']) > 0] or [0])
if not bulan_terisi: sys.exit('Belum ada transaksi bulanan di Excel.')

# ---------------- pos kas lain dari sheet KAS (n) ----------------
BIAYA = {'Biaya Operasional': '501', 'Brosur RAT': '511', 'Administrasi/ATK': '512', 'Konsumsi RAT': '513',
         'Pengolahan Neraca': '514', 'Rapat Pengurus': '515', 'Uang Duduk': '507', 'Door Prize': '516',
         'Souvenir': '508', 'Souvenir RAT': '508', 'Transport Belanja': '517', 'Honor BP & Pembina': '518'}
KEPALA = {'Pembayaran Honor Pengurus': '509', 'Pembayaran THR': '506', 'Inventaris': '111'}
DANA = {'Dana Anggota': '211', 'Dana Pengurus': '212', 'Dana Kesej. Pegawai': '213', 'Dana Pendidikan': '214',
        'Dana Sosial': '215', 'Dana Pemb. Daerah': '216'}
BAKU = {'Saldo Awal Bulan', 'Simpanan Pokok', 'Simpanan Wajib', 'Simpanan Sukarela', 'Pemberian Pinjaman/Piutang',
        'Setoran Angsuran Piutang', 'Jasa Simpan Pinjam', 'J U M L A H'}
peringatan = []

def pos_kas(b):
    hasil = []
    for r in wb[f'KAS ({b})'].iter_rows(min_row=10, max_row=75, values_only=True):
        f_, g_, h_ = rapi(r[5]), rapi(r[6]), rapi(r[7])
        masuk, keluar = r2(r[8]), r2(r[9])
        if not masuk and not keluar: continue
        if rapi(r[3]).replace(' ', '').upper() == 'JUMLAH': continue   # baris total
        label = h_ or g_ or f_
        if f_ in BAKU or g_ in BAKU or label in BAKU or 'SALDO KAS' in h_.upper(): continue
        if f_.startswith('SHU Tahun'):
            if keluar: hasil.append({'jenis': 'alokasi_shu', 'tahun_shu': tahun - 1, 'nominal': keluar})
            continue
        if h_ == 'Dana Cadangan':
            if keluar: peringatan.append(f'KAS ({b}): pengeluaran Dana Cadangan {keluar:,.2f} tidak dimigrasi — catat manual')
            continue  # masuk = bagian alokasi SHU
        if h_ in DANA:
            if keluar: hasil.append({'jenis': 'beban', 'akun': DANA[h_], 'nominal': keluar, 'ket': 'Pengeluaran ' + h_.lower()})
            continue  # masuk = bagian alokasi SHU (sudah dicatat lewat alokasi_shu)
        if f_ in KEPALA:
            hasil.append({'jenis': 'beban', 'akun': KEPALA[f_], 'nominal': keluar, 'ket': f_.replace('Pembayaran ', '')})
            continue
        if 'Penyusutan' in label:
            hasil.append({'jenis': 'penyusutan', 'nominal': keluar or masuk, 'ket': 'Penyusutan inventaris'})
            continue
        if g_ in BIAYA:
            hasil.append({'jenis': 'beban', 'akun': BIAYA[g_], 'nominal': keluar, 'ket': g_})
            continue
        if keluar:
            peringatan.append(f'KAS ({b}): "{label}" tidak dikenal → dicatat Biaya Lain-lain (520) {keluar:,.2f}')
            hasil.append({'jenis': 'beban', 'akun': '520', 'nominal': keluar, 'ket': label or 'Biaya lain-lain'})
        else:
            peringatan.append(f'KAS ({b}): "{label}" pemasukan → dicatat Penghasilan Lain-lain {masuk:,.2f}')
            hasil.append({'jenis': 'masuk', 'nominal': masuk, 'ket': label or 'Penghasilan lain-lain'})
    return hasil

# ---------------- transaksi bulanan ----------------
bulan = []
for b in range(1, bulan_terisi + 1):
    setor = per_anggota(wb[f'Isim({b})'], 11, 3, 11, [4, 5, 6], [12, 13, 14])
    ambil = per_anggota(wb[f'Abl({b})'], 11, 3, 10, [4, 5, 6], [11, 12, 13])
    angs = per_anggota(wb[f'Ang ({b})'], 9, 5, 10, [6, 7], [11, 12])
    pinj = per_anggota(wb[f'Ipin ({b})'], 10, 4, 8, [5], [9])
    simp_akhir = per_anggota(wb[f'Simp ({b})'], 10, 3, 10, [4, 5, 6], [11, 12, 13])
    piut = per_anggota(wb[f'Pinj ({b})'], 9, 3, 10, [4, 5, 6, 7], [11, 12, 13, 14])
    kas_akhir = masuk = keluar = None
    for r in wb[f'KAS ({b})'].iter_rows(min_row=40, max_row=80, values_only=True):
        if r[3] == 'J U M L A H': masuk, keluar = r2(r[8]), r2(r[9])
        if isinstance(r[7], str) and 'SALDO KAS' in r[7]: kas_akhir = r2(r[10])
    bulan.append({'bulan': b,
                  'setoran': {k: v for k, v in setor.items() if any(v)}, 'ambil': {k: v for k, v in ambil.items() if any(v)},
                  'angsuran': {k: v for k, v in angs.items() if any(v)}, 'pinjaman': {k: v[0] for k, v in pinj.items() if v[0]},
                  'kas_lain': pos_kas(b), 'cek_simpanan': simp_akhir, 'cek_piutang': piut,
                  'cek_kas': {'masuk': masuk, 'keluar': keluar, 'saldo_akhir': kas_akhir}})

# ---------------- koreksi label (opsional) ----------------
if arg.koreksi:
    for k in json.load(open(arg.koreksi, encoding='utf8')):
        cocok = [x for x in bulan[k['bulan'] - 1]['kas_lain'] if x.get('akun') == str(k['akun_dari'])
                 and abs(x['nominal'] - float(k['nominal'])) < 0.01]
        if len(cocok) != 1: sys.exit(f'Koreksi tidak cocok dengan tepat satu pos: {k}')
        cocok[0]['akun'] = str(k['akun_ke'])
        cocok[0]['ket'] = k.get('ket', cocok[0]['ket'])

# ---------------- rekonsiliasi ----------------
masalah = []
saldo = {k: list(simp_awal.get(k, [0, 0, 0])) for k in KODE}
piutang = {k: piut_awal.get(k, [0])[0] for k in KODE}
kas = koperasi['kas']
alokasi = {'211': 50, '212': 10, '213': 5, '214': 5, '215': 2.5, '216': 2.5}
saldo_dana = {a: koperasi.get(j, 0) for j, a in (('dana_anggota', '211'), ('dana_pengurus', '212'), ('dana_kesejahteraan', '213'),
                                                ('dana_pendidikan', '214'), ('dana_sosial', '215'), ('dana_pembangunan', '216'))}
for bl in bulan:
    b = bl['bulan']
    for k in KODE:
        s, a = bl['setoran'].get(k, [0, 0, 0]), bl['ambil'].get(k, [0, 0, 0])
        for j in range(3): saldo[k][j] = round(saldo[k][j] + s[j] - a[j], 2)
        exp = bl['cek_simpanan'].get(k, [0, 0, 0])
        for j, nm in enumerate(('pokok', 'wajib', 'sukarela')):
            if abs(saldo[k][j] - exp[j]) > 1:
                masalah.append(f'bulan {b} {KODE[k]["nama"]} {nm}: hitung {saldo[k][j]:,.2f} ≠ Excel {exp[j]:,.2f}')
        ang, pj = bl['angsuran'].get(k, [0, 0]), bl['pinjaman'].get(k, 0)
        piutang[k] = round(piutang[k] - ang[0] + pj, 2)
        cp = bl['cek_piutang'].get(k, [0, 0, 0, 0])
        if abs(piutang[k] - cp[3]) > 1:
            masalah.append(f'bulan {b} {KODE[k]["nama"]} piutang: hitung {piutang[k]:,.2f} ≠ Excel {cp[3]:,.2f}')
        if ang[0] == 0 and ang[1] > 0 and piutang[k] + ang[0] <= 0:
            peringatan.append(f'bulan {b}: {KODE[k]["nama"]} membayar jasa {ang[1]:,.2f} tanpa sisa pinjaman')
        if saldo[k][2] < -1 and (b == 1 or bl['ambil'].get(k, [0, 0, 0])[2]):
            peringatan.append(f'bulan {b}: {KODE[k]["nama"]} saldo sukarela minus {saldo[k][2]:,.2f}')
    m = sum(sum(v) for v in bl['setoran'].values()) + sum(v[0] + v[1] for v in bl['angsuran'].values())
    kk = sum(sum(v) for v in bl['ambil'].values()) + sum(bl['pinjaman'].values())
    for x in bl['kas_lain']:
        if x['jenis'] == 'beban': kk += x['nominal']
        if x['jenis'] == 'masuk': m += x['nominal']
        if x['jenis'] == 'alokasi_shu':
            for a, pr in alokasi.items(): saldo_dana[a] += x['nominal'] * pr / 100
        if x['jenis'] == 'beban' and x['akun'] in saldo_dana:
            saldo_dana[x['akun']] -= x['nominal']
            if saldo_dana[x['akun']] < -1:
                peringatan.append(f'bulan {b}: pengeluaran dana akun {x["akun"]} {x["nominal"]:,.2f} melebihi saldonya — '
                                  f'label di KAS mungkin tertukar (pakai --koreksi)')
    kas = round(kas + m - kk, 2)
    ck = bl['cek_kas']
    if ck['saldo_akhir'] is not None and abs(kas - ck['saldo_akhir']) > 1:
        masalah.append(f'bulan {b} KAS: hitung {kas:,.2f} ≠ Excel {ck["saldo_akhir"]:,.2f}')

print(f'Tahun buku {tahun}: {len(anggota)} anggota, bulan 1–{bulan_terisi}')
print(f'Saldo 31 Des {tahun - 1}: kas {koperasi["kas"]:,.2f}; piutang {sum(v[0] for v in piut_awal.values()):,.2f}')
for x in peringatan: print('  ⚠', x)
if masalah:
    print(f'\n✗ {len(masalah)} selisih dengan Excel — file TIDAK dibuat:')
    for x in masalah[:40]: print('  -', x)
    sys.exit(1)

data = {'sumber': os.path.basename(arg.excel), 'tahun': tahun, 'bulan_terisi': bulan_terisi, 'anggota': anggota,
        'awal': {'simpanan': simp_awal, 'piutang': {k: v[0] for k, v in piut_awal.items() if v[0]}, 'koperasi': koperasi},
        'alokasi_shu_persen': {'cadangan': 25, 'dana_anggota': 50, 'dana_pengurus': 10, 'dana_kesejahteraan': 5,
                               'dana_pendidikan': 5, 'dana_sosial': 2.5, 'dana_pembangunan': 2.5},
        'aturan': {'jasa_pinjaman': arg.jasa_pinjaman, 'jasa_sukarela': arg.jasa_sukarela, 'metode_jasa_tahun_ini': arg.metode_jasa,
                   'metode_jasa_tahun_lalu': 'harian', 'nominal_pokok': arg.pokok, 'nominal_wajib': arg.wajib,
                   'minimal_sukarela': arg.minimal_sukarela},
        'bulan': bulan}
with open(arg.keluar, 'w', encoding='utf8', newline='\n') as f:
    f.write('/**\n * DATA MIGRASI — dibuat oleh tools/ekstrak_excel.py dari "' + os.path.basename(arg.excel) + '".\n'
            ' * Berisi nama & saldo anggota: JANGAN di-commit; hapus dari Apps Script setelah migrasi.\n */\n'
            'const DATA_MIGRASI = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n')
print(f'\n✓ Cocok dengan Excel setiap bulan. Ditulis: {os.path.abspath(arg.keluar)}')
